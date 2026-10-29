"""Export an editable Practical 1 comparison/journal workbook (OOXML).
No benchmark numbers, engine outputs or invented question marks.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from nosql_content import MODELS, CASES, CAP_CASES, SOURCES
from subject_DI04000031 import ROOT

OUT = os.path.join(ROOT, 'preview', 'DI04000031_unit1_practical1_worksheet.xlsx')
PURPLE = '6E28B8'; LIGHT = 'EEE4F8'; BLUE = '1556A1'; PALE = 'EDF5FF'
INK = '252033'; GREY = '675E75'


def setup(ws, title, subtitle, headers, widths):
    ws.sheet_view.showGridLines = False
    n = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    c = ws.cell(1, 1, title); c.font = Font(name='Calibri', size=19, bold=True, color=PURPLE)
    ws.row_dimensions[1].height = 33
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    c = ws.cell(2, 1, subtitle)
    c.font = Font(name='Calibri', size=11, color=GREY)
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[2].height = 34
    for col, label in enumerate(headers, 1):
        c = ws.cell(4, col, label)
        c.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=PURPLE)
        c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[4].height = 34
    for i, w in enumerate(widths, 1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'B5'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if n > 5 else ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:4'
    ws.sheet_properties.tabColor = PURPLE
    ws.oddFooter.center.text = 'DI04000031 · Practical 1 · &P / &N'


def rows(ws, data, height=100, start=5):
    for r, vals in enumerate(data, start):
        for col, val in enumerate(vals, 1):
            c = ws.cell(r, col, val)
            c.font = Font(name='Calibri', size=11, color=INK)
            c.fill = PatternFill('solid', fgColor='F8F5FC' if r % 2 else 'FFFFFF')
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.border = Border(bottom=Side(style='hair', color='D8CDE5'))
        ws.row_dimensions[r].height = height
    ws.print_area = f'A1:{ws.cell(ws.max_row, ws.max_column).coordinate}'


def export(path=OUT):
    wb = Workbook(); wb.remove(wb.active)
    wb.properties.title = 'DI04000031 — Practical 1 Comparative Study'
    wb.properties.subject = 'Introduction to NoSQL, Unit 1'
    wb.properties.creator = 'DiplomaOne'
    wb.properties.description = 'Worked comparisons and an editable student journal; no claimed server executions or performance measurements.'
    ws = wb.create_sheet('Read Me')
    setup(ws, 'Practical 1 · NoSQL Comparative Study',
          'GTU DI04000031 · Computer Semester 4 · English · Unit 1: 8 hours / 13%',
          ['Item', 'Instructions / reference'], [28, 105])
    info = [
        ['Aim', 'Compare Document, Key-Value, Wide-Column and Graph databases: characteristics, advantages and application areas.'],
        ['How to use', 'Read the Four Models and Solved Choices sheets, discuss CAP Cases, then complete Your Journal with your own application and reasoning.'],
        ['Evidence boundary', 'All displayed records and cases are teaching illustrations. This workbook contains no captured MongoDB, Redis, Cassandra or Neo4j runtime output and no benchmarks.'],
        ['Assessment', '0 theory ESE + 0 theory PA + 20 practical PA + 30 practical ESE = 50 total; 3 credits. This worksheet does not assign official per-question marks.'],
        ['Scope', 'Practical 1 is a comparative study. Engine setup and actual operations occur in the later product-specific practicals.'],
        ['Selection', 'Write the real data shape and top reads/writes, state guarantees, choose a starting model, and include two reasons plus one limitation. Well-justified alternatives may be valid.'],
        ['Editable cells', 'The pale-blue cells in Your Journal are for student input. The comparison is qualitative: no automatic score pretends there is one universally correct product.'],
    ] + [[name, url] for name, url in SOURCES]
    rows(ws, info, 55)
    for r in range(12, ws.max_row + 1):
        cell = ws.cell(r, 2)
        if isinstance(cell.value, str) and cell.value.startswith('https://'):
            cell.hyperlink = cell.value; cell.font = Font(name='Calibri', size=11, color=BLUE, underline='single')

    ws = wb.create_sheet('Four Models')
    setup(ws, 'The Four NoSQL Models', 'Typical fits, not universal performance rankings. Product features and guarantees must be evaluated.',
          ['Type', 'Example', 'Representation', 'Typical access', 'Advantage', 'Trade-off', 'Application area', 'Avoid this assumption'],
          [18, 23, 45, 43, 45, 47, 38, 45])
    rows(ws, [[m[k] for k in ('type', 'example', 'shape', 'access', 'advantage', 'tradeoff', 'application', 'not_for')] for m in MODELS], 115)
    ws.auto_filter.ref = 'A4:H8'; ws.freeze_panes = 'C5'

    ws = wb.create_sheet('Solved Choices')
    setup(ws, 'Eight Worked Selection Cases', 'The suggested starting type assumes the stated workload; defend your reasoning and acknowledge trade-offs.',
          ['Case', 'Scenario', 'Main access', 'Starting type', 'Reason', 'Trade-off / caution'],
          [10, 58, 47, 21, 58, 62])
    rows(ws, [[c[k] for k in ('id', 'scenario', 'access', 'choice', 'reason', 'tradeoff')] for c in CASES], 95)
    ws.auto_filter.ref = 'A4:F12'

    ws = wb.create_sheet('CAP Cases')
    setup(ws, 'CAP · Basic Failure Scenarios', 'Analogy and definitions only: no theorem proof, quorum calculation or unconditional product labels.',
          ['Situation', 'Action / policy', 'Meaning'], [37, 67, 74])
    rows(ws, [[c['case'], c['action'], c['result']] for c in CAP_CASES], 94)

    ws = wb.create_sheet('Your Journal')
    setup(ws, 'Your Practical 1 Journal', 'Replace the pale-blue prompts with your own work. No server installation or fabricated timing figures are required.',
          ['Field / model', 'Your response', 'Access / advantage', 'Trade-off / application'], [30, 50, 48, 49])
    prompts = [
        ['Student name / enrollment', '', '', ''],
        ['Date / class / batch', '', '', ''],
        ['Aim', 'Compare all four NoSQL types.', '', ''],
        ['Chosen application', '', '', ''],
        ['Entities / data shape', '', '', ''],
        ['Top read requests', '', '', ''],
        ['Top write requests', '', '', ''],
        ['Consistency / durability needs', '', '', ''],
        ['Chosen starting type', '', '', ''],
        ['Reason 1', '', '', ''],
        ['Reason 2', '', '', ''],
        ['Limitation / alternative', '', '', ''],
        ['Document comparison', '', '', ''],
        ['Key-Value comparison', '', '', ''],
        ['Wide-Column comparison', '', '', ''],
        ['Graph comparison', '', '', ''],
        ['Model sketch / representative record', '', '', ''],
        ['CAP analogy during a partition', '', '', ''],
        ['Deferral / staleness policy', '', '', ''],
        ['Conclusion', '', '', ''],
        ['Sources', '', '', ''],
        ['Evidence boundary', 'Study observations / illustrative records; list actual runtime evidence only if you really ran it.', '', ''],
    ]
    rows(ws, prompts, 58)
    for r in range(5, ws.max_row + 1):
        for col in range(2, 5):
            ws.cell(r, col).fill = PatternFill('solid', fgColor=PALE)
            ws.cell(r, col).font = Font(name='Calibri', size=11, color=BLUE)
    for r in (21, 22): ws.row_dimensions[r].height = 105
    dv = DataValidation(type='list', formula1='"Document,Key-Value,Wide-Column,Graph,Relational"', allow_blank=True)
    dv.errorTitle = 'Choose a starting type'; dv.error = 'Select a model from the list; justify it in the next rows.'
    dv.showErrorMessage = True; dv.showInputMessage = True
    dv.promptTitle = 'Starting model'; dv.prompt = 'Choose from workload needs, not a brand name.'
    ws.add_data_validation(dv); dv.add(ws['B13'])
    ws.row_breaks.append(Break(id=16))
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    print('Practical 1 workbook:', path, os.path.getsize(path), 'bytes /', len(wb.worksheets), 'sheets')
    return path


if __name__ == '__main__': export()
