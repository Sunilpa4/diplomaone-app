"""Unit 1 comparative-study regression tests — not NoSQL engine runs.

Checks curriculum, fixtures/reasoning, HTML/SVGs, worksheet, preview JS and
export integration. Actual MongoDB/Cassandra/Neo4j/Redis execution belongs
to later units and is intentionally not claimed here.
"""
import collections
import glob
import hashlib
import html
from html.parser import HTMLParser
import io
import json
import os
from pathlib import Path
import re
import subprocess
import tempfile
import unittest
import xml.etree.ElementTree as ET

from PIL import Image, ImageDraw, ImageFont
import resvg_py
from openpyxl import load_workbook

from subject_DI04000031 import ROOT, SUBJ_DIR, DIAG_DIR, TABS
from nosql_content import DOCUMENTS, CACHE, SENSOR_ROWS, GRAPH_NODES, GRAPH_EDGES, MODELS, CASES, CAP_CASES, SOURCES
from content_DI04000031_u1_notes import NOTES, SN
from content_DI04000031_u1_extras import PRACTICE_QS
from diagrams_DI04000031_u1 import TITLES
from preview_builder import build_preview
from export_DI04000031_u1_worksheet import OUT as WORKBOOK

PREVIEW = os.path.join(ROOT, 'preview', 'DI04000031_unit1_preview.html')


def load(name):
    return json.loads(Path(SUBJ_DIR, name).read_text(encoding='utf-8'))


def strings(obj):
    if isinstance(obj, str): yield obj
    elif isinstance(obj, dict):
        for v in obj.values(): yield from strings(v)
    elif isinstance(obj, list):
        for v in obj: yield from strings(v)


class Fragment(HTMLParser):
    allowed = {'p', 'b', 'strong', 'i', 'em', 'code', 'pre', 'ul', 'ol', 'li', 'br',
               'table', 'tr', 'th', 'td', 'span', 'div', 'h2', 'h3', 'h4', 'a', 'sup', 'sub'}
    def __init__(self):
        super().__init__(convert_charrefs=True); self.stack = []
    def handle_starttag(self, tag, attrs):
        assert tag in self.allowed, 'Unexpected unescaped HTML tag: ' + tag
        assert not any(k.lower().startswith('on') for k, v in attrs), 'Event attribute'
        if tag == 'a':
            assert dict(attrs).get('href', '').startswith('https://'), 'Unsafe source link'
        if tag != 'br': self.stack.append(tag)
    def handle_endtag(self, tag):
        assert self.stack and self.stack[-1] == tag, ('Unbalanced HTML', tag, self.stack)
        self.stack.pop()
    def finish(self, text):
        self.feed(text); self.close(); assert not self.stack, self.stack


def execute_preview(path):
    content = Path(path).read_text(encoding='utf-8')
    scripts = re.findall(r'<script>(.*?)</script>', content, re.S)
    assert len(scripts) == 1
    script = scripts[0]
    syntax = subprocess.run(['node', '--check'], input=script, text=True, capture_output=True)
    assert syntax.returncode == 0, syntax.stderr
    prefix = """
const elements = {};
const clickHandlers = [];
const document = {
 getElementById(id) {
  return elements[id] ||= {style:{}, innerHTML:'', textContent:'', classList:{add(){},remove(){}}};
 },
 querySelectorAll(){return [];},
 addEventListener(event, fn){ if(event==='click') clickHandlers.push(fn); }
};
const window = {scrollTo(){}};
"""
    suffix = r"""
const views = {};
for (const [key, fn] of Object.entries(VIEWS)) {
 if (typeof fn !== 'function') throw new Error('Unknown view ' + key);
 S.tab=key; render(); views[key]=elements.view.innerHTML;
 if (/>\s*undefined\s*</.test(views[key])) throw new Error('Undefined content in ' + key);
}
process.stdout.write(JSON.stringify({views, languageHidden:elements.lang.style.display,
 tabs:elements.tabs.innerHTML, neutral:hl('a < b && c > d'), handlers:clickHandlers.length}));
"""
    r = subprocess.run(['node'], input=prefix + script + suffix, text=True, capture_output=True)
    assert r.returncode == 0, r.stderr
    return json.loads(r.stdout)


class Unit1Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.u = load('unit1.json'); cls.s = load('subject.json')

    def test_01_official_subject_metadata(self):
        s = self.s
        self.assertEqual((s['semester'], s['syllabus_year'], s['type']), (4, '2025-26', 'practical'))
        self.assertEqual(s['exam'], {'theory_ese': 0, 'theory_pa': 0, 'practical_pa': 20,
                                     'practical_ese': 30, 'total': 50, 'credits': 3})
        self.assertEqual([u['hours'] for u in s['units']], [8, 18, 12, 12, 10])
        self.assertEqual([u['weightage'] for u in s['units']], [13, 30, 20, 20, 17])
        self.assertEqual(sum(u['hours'] for u in s['units']), 60)
        self.assertEqual(sum(u['weightage'] for u in s['units']), 100)
        self.assertEqual([u['file'] for u in s['units']], ['unit1.json', None, None, None, None])
        self.assertEqual(len(s['practicals']), 9)

    def test_02_unit_coverage_and_counts(self):
        u = self.u
        self.assertEqual((u['unit'], u['hours'], u['weightage']), (1, 8, 13))
        self.assertEqual(sum(len(s['items']) for s in NOTES), 14)
        self.assertEqual(sum(len(s['items']) for s in u['notes']), 30)
        self.assertEqual(len(u['notes'][-1]['items']), 16)
        self.assertEqual(sum(len(g['items']) for g in u['syntax']), 12)
        self.assertEqual(sum(len(g['items']) for g in u['viva']), 48)
        self.assertEqual(len(u['diagrams']), 10)
        self.assertEqual(u['practical_numbers'], [1])
        self.assertEqual(u['worked_selection_cases'], 8)
        self.assertTrue(all(t['body']['en'].strip() for sec in u['notes'] for t in sec['items']))

    def test_03_english_only_and_no_key_stars(self):
        def check(o):
            if isinstance(o, dict):
                self.assertNotIn('gu', o)
                self.assertNotIn('key', o)
                for v in o.values(): check(v)
            elif isinstance(o, list):
                for v in o: check(v)
            elif isinstance(o, str):
                self.assertIsNone(re.search(r'[\u0a80-\u0aff]', o))
        check(self.u); check(self.s)
        self.assertEqual(self.s['languages'], ['en'])

    def test_04_honest_conceptual_scope_and_practical_assessment(self):
        u = self.u
        self.assertEqual(u['content_kind'], 'conceptual_comparison')
        self.assertEqual(u['tabs'], ['notes', 'diagrams', 'syntax', 'viva'])
        self.assertEqual(u['tabs'], TABS)
        self.assertEqual(u['exam_questions'], [])
        self.assertNotIn('programs', u)
        self.assertNotIn('out', u)
        self.assertIn('not database-server output', u['intro_text']['en'])
        self.assertFalse(any('marks' in q for q in PRACTICE_QS))
        for q, item in zip(PRACTICE_QS, u['notes'][-1]['items']):
            self.assertEqual(q['q'], item['h']); self.assertEqual(q['a'], item['body'])

    def test_05_document_json_models(self):
        self.assertEqual(json.loads(SN['doc_sensor']), DOCUMENTS[0])
        self.assertEqual(json.loads(SN['doc_book']), DOCUMENTS[1])
        self.assertEqual(len({d['_id'] for d in DOCUMENTS}), 2)
        self.assertNotEqual(set(DOCUMENTS[0]['attributes']), set(DOCUMENTS[1]['attributes']))
        self.assertIs(DOCUMENTS[0]['attributes']['battery_powered'], True)
        self.assertEqual(DOCUMENTS[1]['attributes']['pages'], 320)

    def test_06_known_key_cache_model(self):
        self.assertRegex(CACHE['key'], r'^recs:U\d+:v\d+$')
        self.assertEqual(CACHE['ttl_seconds'], 300)
        ids = {d['_id'] for d in DOCUMENTS}
        self.assertTrue(set(CACHE['value']) <= ids)
        self.assertIn(CACHE['key'], SN['kv'])
        self.assertIn('miss', SN['kv'])

    def test_07_sensor_partition_fixture(self):
        groups = collections.defaultdict(list)
        unique = set()
        for r in SENSOR_ROWS:
            key = r['sensor_id'], r['day']
            groups[key].append(r)
            self.assertNotIn(key + (r['ts'],), unique); unique.add(key + (r['ts'],))
            self.assertIsInstance(r['reading_c'], float)
        self.assertEqual(len(groups), 2)
        self.assertEqual([r['ts'] for r in groups[('S01', '2026-10-29')]], ['09:00', '09:01'])
        self.assertEqual(len(groups[('S02', '2026-10-29')]), 1)

    def test_08_graph_fixture_and_reasoned_candidate(self):
        ids = {n['id'] for n in GRAPH_NODES}
        likes = collections.defaultdict(set)
        for a, rel, b in GRAPH_EDGES:
            self.assertTrue({a, b} <= ids); self.assertEqual(rel, 'LIKES'); likes[a].add(b)
        # Local fixture reasoning, explicitly NOT a Neo4j query or model-training run.
        own = likes['U42']
        peers = [u for u in likes if u != 'U42' and likes[u] & own]
        candidates = set().union(*(likes[u] for u in peers)) - own
        self.assertEqual(peers, ['U77']); self.assertEqual(candidates, {'P102'})

    def test_09_comparison_and_selection_reasoning(self):
        self.assertEqual({m['type'] for m in MODELS}, {'Document', 'Key-Value', 'Wide-Column', 'Graph'})
        self.assertEqual([c['id'] for c in CASES], ['S' + str(n) for n in range(1, 9)])
        self.assertEqual([c['choice'] for c in CASES[:5]], ['Document', 'Key-Value', 'Wide-Column', 'Graph', 'Relational'])
        self.assertTrue(all(c['reason'] and c['access'] and c['tradeoff'] for c in CASES))
        self.assertEqual(CASES[5]['choice'], 'Relational')

    def test_10_cap_accuracy_guards(self):
        cap = ' '.join(t['body']['en'] for t in self.u['notes'][1]['items'])
        self.assertIn('during a network partition', cap.lower())
        self.assertIn('linearizability', cap)
        self.assertIn('does not set a fixed response-time bound', cap)
        self.assertIn('not the C of ACID', cap)
        self.assertIn('sold out', cap)
        self.assertIn('CP does not mean one office can always proceed', cap)
        self.assertEqual(len(CAP_CASES), 4)
        self.assertIn('can both be provided', CAP_CASES[0]['result'])
        self.assertIn('not automatic', CAP_CASES[-1]['result'])

    def test_11_sources_and_syllabus_link(self):
        self.assertEqual(len(self.u['sources']), 6)
        self.assertEqual(self.u['sources'][0]['url'], self.s['syllabus_url'])
        for name, url in SOURCES:
            self.assertTrue(url.startswith('https://'))
            self.assertIn(url, self.u['notes'][4]['items'][1]['body']['en'])
        self.assertIn('CO1', self.u['notes'][0]['items'][0]['body']['en'])

    def test_12_safe_balanced_note_and_answer_html(self):
        for sec in self.u['notes']:
            for item in sec['items']:
                Fragment().finish(item['body']['en'])
                self.assertNotIn('@@', item['body']['en'])
        for group in self.u['viva']:
            for item in group['items']: Fragment().finish(item['a']['en'])
        for group in self.u['syntax']:
            for item in group['items']: Fragment().finish(item['terms']['en'])

    def test_13_svg_ids_links_and_self_containment(self):
        ds = self.u['diagrams']
        self.assertEqual({d['id'] for d in ds}, set(TITLES))
        used = {i['diagram'] for s in self.u['notes'] for i in s['items'] if 'diagram' in i}
        self.assertEqual(used, set(TITLES))
        for d in ds:
            path = Path(SUBJ_DIR, d['file'])
            raw = path.read_text(); root = ET.fromstring(raw)
            self.assertEqual(root.attrib['viewBox'].split()[2], '360')
            for bad in ('<marker', '<image', '<foreignObject', '<script', '<style', 'href=', 'url('):
                self.assertNotIn(bad, raw)
            self.assertEqual(d['title']['en'], TITLES[d['id']])

    def test_14_every_svg_rasterizes_and_contact_sheets(self):
        checks = Path(ROOT, 'preview', 'nosql_u1_checks'); checks.mkdir(exist_ok=True)
        thumbs = []
        for id_ in TITLES:
            b = resvg_py.svg_to_bytes(svg_path=str(Path(DIAG_DIR, id_ + '.svg')), zoom=1.4,
                font_dirs=['/usr/share/fonts/truetype/dejavu'], font_family='DejaVu Sans',
                sans_serif_family='DejaVu Sans', skip_system_fonts=True)
            image = Image.open(io.BytesIO(bytes(b))).convert('RGB')
            self.assertGreater(image.height, 400)
            image.save(checks / (id_ + '.png')); thumbs.append((id_, image))
        for batch, label in [(thumbs[:5], 'a'), (thumbs[5:], 'b')]:
            width = 530; height = 735
            sheet = Image.new('RGB', (width * 2, height * 3), '#EAE4F1')
            draw = ImageDraw.Draw(sheet)
            for i, (id_, image) in enumerate(batch):
                x = (i % 2) * width + 12; y = (i // 2) * height + 30
                draw.text((x, y - 22), id_, fill='#4C1C85')
                sheet.paste(image, (x, y))
            sheet.save(Path(ROOT, 'preview', f'DI04000031_unit1_diagrams_{label}.png'))

    def test_15_svg_text_stays_inside_viewport(self):
        for id_ in TITLES:
            root = ET.parse(Path(DIAG_DIR, id_ + '.svg')).getroot()
            _, _, width, height = map(float, root.attrib['viewBox'].split())
            for node in root.iter('{http://www.w3.org/2000/svg}text'):
                size = float(node.get('font-size', '12'))
                fontname = 'DejaVuSans-Bold.ttf' if node.get('font-weight') == 'bold' else 'DejaVuSans.ttf'
                font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/' + fontname, round(size * 10))
                tw = font.getlength(''.join(node.itertext())) / 10
                x = float(node.get('x')); y = float(node.get('y'))
                anchor = node.get('text-anchor', 'start')
                left = x - tw / 2 if anchor == 'middle' else x - tw if anchor == 'end' else x
                self.assertGreaterEqual(left, 1, (id_, node.text))
                self.assertLessEqual(left + tw, width - 1, (id_, node.text, left + tw))
                self.assertGreater(y - size, -1); self.assertLess(y + size / 3, height + 1)

    def test_16_workbook_structure_answers_and_editable_template(self):
        wb = load_workbook(WORKBOOK)
        self.assertEqual(wb.sheetnames, ['Read Me', 'Four Models', 'Solved Choices', 'CAP Cases', 'Your Journal'])
        self.assertEqual(wb['Four Models'].max_row, 8)
        self.assertEqual(wb['Solved Choices'].max_row, 12)
        self.assertEqual(wb['CAP Cases'].max_row, 8)
        self.assertEqual([wb['Four Models'].cell(r, 1).value for r in range(5, 9)], [m['type'] for m in MODELS])
        self.assertEqual(wb['Solved Choices']['D9'].value, 'Relational')
        self.assertEqual(wb['Your Journal']['A13'].value, 'Chosen starting type')
        self.assertIsNone(wb['Your Journal']['B13'].value)
        self.assertIn('B13', str(wb['Your Journal'].data_validations.dataValidation[0].sqref))
        self.assertFalse(any(c.data_type == 'f' for ws in wb for row in ws for c in row))
        self.assertFalse(wb._external_links)
        self.assertTrue(all(ws.freeze_panes for ws in wb))
        wb.close()

    def test_17_preview_javascript_views_intro_and_language(self):
        text = Path(PREVIEW).read_text()
        self.assertNotRegex(text, r'<(?:script|link|img)[^>]+(?:src|href)=')
        result = execute_preview(PREVIEW)
        self.assertEqual(set(result['views']), {'notes', 'diagrams', 'syntax', 'viva'})
        self.assertEqual(result['languageHidden'], 'none')
        self.assertEqual(result['handlers'], 1)
        self.assertNotIn('data-tab="exam"', result['tabs'])
        self.assertIn(self.u['intro_text']['en'], result['views']['notes'])
        self.assertNotIn('solved C programs', result['views']['notes'])
        self.assertEqual(result['neutral'], 'a &lt; b &amp;&amp; c &gt; d')
        self.assertEqual(result['views']['diagrams'].count('class="dcard"'), 10)

    def test_18_preview_builder_legacy_regression(self):
        specs = [('DI04000011', 'sem4', 1), ('DI04007011', 'sem4', 5)]
        for path in glob.glob(os.path.join(ROOT, 'content', 'computer', 'sem*', '*', 'subject.json')):
            subject = json.loads(Path(path).read_text())
            if subject.get('code_lang') == 'python':
                specs.append((subject['code'], Path(path).parts[-3], 1)); break
        self.assertGreaterEqual(len(specs), 3)
        with tempfile.TemporaryDirectory() as tmp:
            for code_, sem, unit_no in specs:
                directory = Path(ROOT, 'content', 'computer', sem, code_)
                unit = json.loads((directory / f'unit{unit_no}.json').read_text())
                subject = json.loads((directory / 'subject.json').read_text())
                path = os.path.join(tmp, code_ + '.html')
                build_preview(unit, subject, str(directory / 'diagrams'), path)
                result = execute_preview(path)
                self.assertIn('notes', result['views'])
                self.assertTrue(result['views']['notes'])

    def test_19_rebuild_is_deterministic_for_published_sources(self):
        files = [Path(SUBJ_DIR, 'unit1.json'), Path(SUBJ_DIR, 'subject.json'), Path(PREVIEW)]
        files += [Path(DIAG_DIR, id_ + '.svg') for id_ in TITLES]
        before = {str(p): p.read_bytes() for p in files}
        r = subprocess.run(['python3', 'tools/build_DI04000031_u1.py'], cwd=ROOT, text=True, capture_output=True)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertEqual(before, {str(p): p.read_bytes() for p in files})

    def test_20_export_catalog_hashes_and_partial_availability(self):
        with tempfile.TemporaryDirectory() as tmp:
            r = subprocess.run(['python3', 'tools/export_app_content.py', '--out', tmp], cwd=ROOT, text=True, capture_output=True)
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            subject = json.loads(Path(tmp, 'subjects', 'DI04000031.json').read_text())
            unit = json.loads(Path(tmp, 'units', 'DI04000031_u1.json').read_text())
            catalog = json.loads(Path(tmp, 'catalog.json').read_text())
            self.assertEqual(subject['available_units'], [1])
            self.assertEqual(len(subject['units']), 5)
            self.assertEqual(catalog['subjects']['DI04000031']['status'], 'available')
            self.assertEqual(unit['tabs'], TABS)
            self.assertEqual(set(unit['svgs']), set(TITLES))
            for rel in ('subjects/DI04000031.json', 'units/DI04000031_u1.json'):
                self.assertEqual(catalog['files'][rel], hashlib.sha1(Path(tmp, rel).read_bytes()).hexdigest()[:12])
            self.assertFalse(Path(tmp, 'units', 'DI04000031_u2.json').exists())


if __name__ == '__main__': unittest.main(verbosity=2)
